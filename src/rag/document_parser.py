"""文档解析器模块"""

import asyncio
import os
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Any, Dict, Optional

# 文档解析库
try:
    import pypdf

    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    from docx import Document as DocxDocument

    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import openpyxl

    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

import logging

from ..exceptions import DocumentProcessingException
from ..utils.file_utils import get_file_extension, is_supported_file

logger = logging.getLogger("document_parser")


class DocumentParser:
    """文档解析器，支持多种文档格式"""

    def __init__(self):
        self.executor = ThreadPoolExecutor(max_workers=4)

    async def parse_document(self, file_path: str) -> Dict[str, Any]:
        """解析文档并提取内容

        Args:
            file_path: 文档文件路径

        Returns:
            包含文档内容和元数据的字典

        Raises:
            DocumentProcessingException: 文档处理失败时抛出
        """
        if not os.path.exists(file_path):
            raise DocumentProcessingException(f"File not found: {file_path}")

        if not is_supported_file(file_path):
            raise DocumentProcessingException(f"Unsupported file type: {file_path}")

        try:
            # 在线程池中执行文档解析
            loop = asyncio.get_event_loop()
            result = await loop.run_in_executor(
                self.executor, self._parse_document_sync, file_path
            )

            logger.info(f"Successfully parsed document: {file_path}")
            return result

        except Exception as e:
            logger.error(f"Failed to parse document {file_path}: {e}")
            raise DocumentProcessingException(f"Failed to parse document: {e}")

    def _parse_document_sync(self, file_path: str) -> Dict[str, Any]:
        """同步解析文档"""
        extension = get_file_extension(file_path)
        file_name = Path(file_path).name
        file_size = os.path.getsize(file_path)

        # 基础元数据
        metadata = {
            "file_name": file_name,
            "file_path": file_path,
            "file_size": file_size,
            "file_type": extension,
        }

        # 根据文件类型解析内容
        try:
            if extension == ".pdf":
                content = (
                    self._parse_pdf(file_path)
                    if PDF_AVAILABLE
                    else self._parse_fallback(file_path)
                )
            elif extension in [".docx", ".doc"]:
                content = (
                    self._parse_docx(file_path)
                    if DOCX_AVAILABLE
                    else self._parse_fallback(file_path)
                )
            elif extension in [".xlsx", ".xls"]:
                content = (
                    self._parse_excel(file_path)
                    if EXCEL_AVAILABLE
                    else self._parse_fallback(file_path)
                )
            elif extension in [
                ".txt",
                ".md",
                ".py",
                ".js",
                ".html",
                ".css",
                ".json",
                ".xml",
                ".csv",
            ]:
                content = self._parse_text(file_path)
            else:
                # 对于不支持的文件类型，尝试作为文本文件解析
                content = self._parse_fallback(file_path)
        except Exception as e:
            logger.warning(f"主解析器失败，尝试降级解析: {e}")
            content = self._parse_fallback(file_path)

        return {
            "content": content,
            "metadata": metadata,
            "word_count": len(content.split()) if content else 0,
            "char_count": len(content) if content else 0,
        }

    def _parse_pdf(self, file_path: str) -> str:
        """解析PDF文件"""
        if not PDF_AVAILABLE:
            raise DocumentProcessingException(
                "pypdf not installed, cannot parse PDF files"
            )

        try:
            with open(file_path, "rb") as file:
                pdf_reader = pypdf.PdfReader(file)
                text_content = []

                for page_num, page in enumerate(pdf_reader.pages):
                    try:
                        page_text = page.extract_text()
                        if page_text.strip():
                            text_content.append(
                                f"--- Page {page_num + 1} ---\n{page_text}"
                            )
                    except Exception as e:
                        logger.warning(
                            f"Failed to extract text from page {page_num + 1}: {e}"
                        )
                        continue

                return "\n\n".join(text_content)

        except Exception as e:
            raise DocumentProcessingException(f"Failed to parse PDF: {e}")

    def _parse_docx(self, file_path: str) -> str:
        """解析Word文档"""
        if not DOCX_AVAILABLE:
            raise DocumentProcessingException(
                "python-docx not installed, cannot parse DOCX files"
            )

        try:
            doc = DocxDocument(file_path)
            paragraphs = []

            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    paragraphs.append(text)

            # 处理表格
            for table in doc.tables:
                table_text = self._extract_table_text(table)
                if table_text:
                    paragraphs.append(f"--- Table ---\n{table_text}")

            return "\n\n".join(paragraphs)

        except Exception as e:
            raise DocumentProcessingException(f"Failed to parse DOCX: {e}")

    def _parse_excel(self, file_path: str) -> str:
        """解析Excel文件"""
        if not EXCEL_AVAILABLE:
            raise DocumentProcessingException(
                "openpyxl not installed, cannot parse Excel files"
            )

        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            sheets_content = []

            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                sheet_text = self._extract_sheet_text(sheet, sheet_name)
                if sheet_text:
                    sheets_content.append(sheet_text)

            return "\n\n".join(sheets_content)

        except Exception as e:
            raise DocumentProcessingException(f"Failed to parse Excel: {e}")

    def _parse_text(self, file_path: str) -> str:
        """解析纯文本文件"""
        try:
            # 尝试多种编码
            encodings = ["utf-8", "gbk", "gb2312", "latin-1"]

            for encoding in encodings:
                try:
                    with open(file_path, "r", encoding=encoding) as file:
                        content = file.read()
                        logger.info(f"成功使用 {encoding} 编码解析文本文件")
                        return content
                except UnicodeDecodeError:
                    continue

            # 如果所有编码都失败，使用错误处理
            with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
                content = file.read()
                logger.warning("使用UTF-8编码忽略错误解析文本文件")
                return content

        except Exception as e:
            raise DocumentProcessingException(f"Failed to parse text file: {e}")

    def _parse_fallback(self, file_path: str) -> str:
        """降级解析 - 当其他解析器不可用时使用"""
        try:
            # 对于不支持的文件类型，尝试作为文本文件读取
            return self._parse_text(file_path)
        except Exception as e:
            logger.warning(f"降级解析也失败: {e}")
            return f"无法解析文件内容: {os.path.basename(file_path)}"

    def _extract_table_text(self, table) -> str:
        """从Word表格中提取文本"""
        rows = []
        for row in table.rows:
            cells = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                cells.append(cell_text if cell_text else "")
            if any(cells):  # 只添加非空行
                rows.append(" | ".join(cells))

        return "\n".join(rows)

    def _extract_sheet_text(self, sheet, sheet_name: str) -> str:
        """从Excel工作表中提取文本"""
        rows = []
        header_added = False

        for row in sheet.iter_rows(values_only=True):
            # 过滤空行
            row_values = [str(cell) if cell is not None else "" for cell in row]
            if any(val.strip() for val in row_values):
                if not header_added:
                    rows.append(f"--- Sheet: {sheet_name} ---")
                    header_added = True
                rows.append(" | ".join(row_values))

        return "\n".join(rows)

    async def parse_multiple_documents(
        self, file_paths: list[str]
    ) -> Dict[str, Dict[str, Any]]:
        """批量解析多个文档

        Args:
            file_paths: 文档文件路径列表

        Returns:
            文件路径到解析结果的映射
        """
        results = {}

        # 并发解析文档
        tasks = []
        for file_path in file_paths:
            task = asyncio.create_task(self._parse_single_document(file_path))
            tasks.append((file_path, task))

        for file_path, task in tasks:
            try:
                result = await task
                results[file_path] = result
            except Exception as e:
                logger.error(f"Failed to parse {file_path}: {e}")
                results[file_path] = {
                    "error": str(e),
                    "content": "",
                    "metadata": {"file_path": file_path, "error": True},
                }

        return results

    async def _parse_single_document(self, file_path: str) -> Dict[str, Any]:
        """解析单个文档的包装方法"""
        try:
            return await self.parse_document(file_path)
        except Exception as e:
            raise DocumentProcessingException(f"Failed to parse {file_path}: {e}")

    def __del__(self):
        """清理资源"""
        if hasattr(self, "executor"):
            self.executor.shutdown(wait=False)
